#!/usr/bin/env python
"""
Build a downloadable Word (.docx) bug report with steps to reproduce and
screenshots.

    python scripts/build_bug_report_docx.py

Output: docs/FuseTwo_Bug_Report.docx
"""

import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "FuseTwo_Bug_Report.docx")

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x60, 0x60, 0x60)
RED = RGBColor(0x9C, 0x00, 0x06)
AMBER = RGBColor(0x9C, 0x65, 0x00)


def sev_color(sev):
    return {"Critical": RED, "High": RED, "Medium": AMBER, "Low": GREY}.get(sev, GREY)


# id, severity, title, area, environment, steps[], expected, actual, evidence, screenshot
BUGS = [
    (
        "BUG-1", "High",
        "Programs API returns HTTP 500 across the whole app",
        "Advertiser platform — Programs API",
        "dev (advertiser 1062551)",
        [
            "Sign in to the advertiser platform.",
            "Open any of: Dashboard, Programs, Campaigns, Creatives, Coupons & Offers, "
            "Products, Reports > Sales Detailed, Reports > Clicks Detailed.",
            "Watch the network calls (or the browser console).",
        ],
        "GET /api/{advertiserId}/programs returns 200 with the program list (or an "
        "empty list), and the page shows no error.",
        "GET /api/1062551/programs returns HTTP 500 "
        '{"errorMessage":"Oops! Something went wrong. Please try again in a few moments."} '
        "on every one of those pages, and the browser console logs an error. Reproduced "
        "consistently on repeated loads.",
        "Response body: 500 {\"result\":null,\"errorMessage\":\"Oops! Something went "
        "wrong...\"}. The request is authenticated (valid Bearer token), so this is a "
        "server-side crash, not an auth failure. Note: correlates with the account not "
        "being Approved (currently Initial Setup) — but the API should degrade "
        "gracefully, not 500.",
        "docs/bug-screens/programs.png",
    ),
    (
        "BUG-2", "High",
        "Status-change notification e-mails are not delivered",
        "Management — Advertiser status change",
        "dev",
        [
            "Open an advertiser in the management portal.",
            "Actions > Change Status, pick a new status, submit.",
            "In the 'Advertiser Email Notification' dialog, choose a template and click Send.",
            "Check the advertiser's inbox.",
        ],
        "The chosen template e-mail is delivered to the advertiser.",
        "Send reports success and the dialog closes, but no e-mail arrives. Across many "
        "status changes plus an explicit Send, zero notifications were delivered within "
        "3 minutes, while the signup 'Verify Your Email' consistently arrives within a "
        "minute (so the mail path itself works).",
        "Verified by reading the real QA mailbox (Outlook) after each Send.",
        None,
    ),
    (
        "BUG-3", "High (Security)",
        "Signup accepts a wrong CAPTCHA",
        "Advertiser signup — step 1",
        "dev  (relates to #681)",
        [
            "Open /signup.",
            "Fill step 1 with valid details but type a deliberately wrong CAPTCHA.",
            "Click Continue.",
        ],
        "Registration is blocked and stays on step 1 with a CAPTCHA error.",
        "Registration advances to step 2 (Company Information) despite the wrong CAPTCHA. "
        "Ticket #681 is marked Done on the board but still reproduces on dev.",
        "Automated: e2e/Onboarding/testSignup.spec.ts 'a wrong CAPTCHA is rejected'.",
        None,
    ),
    (
        "BUG-4", "Medium",
        "$50 vs $100 monthly minimum mismatch",
        "Advertiser Collect Payment page vs Management Pricing Schedule",
        "dev  (#762)",
        [
            "Move an advertiser to 'Collect Payment' and sign in as them, OR open the "
            "Collect Payment / Pricing Schedule page.",
            "Compare the monthly minimum shown to the advertiser with the management "
            "Pricing Schedule.",
        ],
        "The monthly minimum is the same in both places.",
        "The advertiser-facing Collect Payment page shows a $50 monthly minimum while the "
        "management Pricing Schedule shows $100. Customer-facing discrepancy.",
        "Screenshot from the onboarding stage capture.",
        "docs/stage-screens/2-CollectPayment.png",
    ),
    (
        "BUG-5", "Low",
        "Email confirmation page shows success for an invalid token",
        "Advertiser e-mail confirmation",
        "dev",
        [
            "Register a new advertiser but do not open the verification e-mail.",
            "Build the confirmation URL yourself: /emailconfirmation?userref=<base64(email)>.",
            "Open it.",
        ],
        "An invalid/forged token is rejected (no success message).",
        "The page shows 'Your email has been verified!' even for a forged token. Sign-in "
        "stays blocked (the account is not actually verified), so it is not an auth "
        "bypass — but the success message on an invalid token is misleading.",
        "Scripted probe; base64(email) is the entire token, with no signature or expiry.",
        None,
    ),
    (
        "BUG-6", "Low",
        "'Required field*' helper appears after valid input",
        "Advertiser signup — step 1",
        "dev  (#686)",
        [
            "Open /signup.",
            "Type a valid value into a field (e.g. First Name).",
            "Observe the helper text.",
        ],
        "No 'Required field*' helper is shown for a field with valid input.",
        "The 'Required field*' helper appears after a valid value is entered and does not "
        "hide.",
        "Automated: e2e/Onboarding/testSignup.spec.ts (#686).",
        None,
    ),
]


def main():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    r = title.add_run("FuseTwo — Bug Report")
    r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = NAVY
    sub = doc.add_paragraph()
    s = sub.add_run(
        "Defects found during QA automation. Each has steps to reproduce, expected vs "
        "actual, evidence and (where captured) a screenshot. Prepared by Syed Shah, QA. "
        "As of 28 August 2026, dev environment."
    )
    s.font.size = Pt(10.5); s.italic = True; s.font.color.rgb = GREY

    # summary table
    doc.add_paragraph()
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    for i, h in enumerate(["ID", "Severity", "Title", "Area"]):
        c = tbl.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.font.bold = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "2E5C8A")
        c._tc.get_or_add_tcPr().append(shd)
    for bug in BUGS:
        cells = tbl.add_row().cells
        vals = [bug[0], bug[1], bug[2], bug[3]]
        for i, v in enumerate(vals):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(v)
            run.font.size = Pt(9)
            if i == 1:
                run.font.color.rgb = sev_color(bug[1].split()[0]); run.font.bold = True

    # detail per bug
    for bug in BUGS:
        bug_id, sev, title_t, area, envr, steps, expected, actual, evidence, shot = bug
        doc.add_paragraph()
        h = doc.add_heading(level=1)
        run = h.add_run(f"{bug_id}: {title_t}")
        run.font.color.rgb = NAVY

        meta = doc.add_paragraph()
        m = meta.add_run(f"Severity: {sev}    |    Area: {area}    |    Environment: {envr}")
        m.font.size = Pt(9.5); m.italic = True; m.font.color.rgb = GREY

        doc.add_paragraph().add_run("Steps to reproduce").bold = True
        for i, step in enumerate(steps, 1):
            p = doc.add_paragraph(style="List Number")
            p.add_run(step).font.size = Pt(10.5)

        p = doc.add_paragraph(); p.add_run("Expected: ").bold = True; p.add_run(expected).font.size = Pt(10.5)
        p = doc.add_paragraph(); rr = p.add_run("Actual: "); rr.bold = True; rr.font.color.rgb = RED
        p.add_run(actual).font.size = Pt(10.5)
        p = doc.add_paragraph(); p.add_run("Evidence: ").bold = True; p.add_run(evidence).font.size = Pt(10.5)

        if shot:
            path = os.path.join(ROOT, shot)
            if os.path.exists(path):
                cap = doc.add_paragraph(); cr = cap.add_run(f"Screenshot: {os.path.basename(shot)}")
                cr.font.size = Pt(9); cr.italic = True; cr.font.color.rgb = GREY
                doc.add_picture(path, width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
