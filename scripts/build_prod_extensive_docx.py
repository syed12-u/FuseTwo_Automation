#!/usr/bin/env python
"""Downloadable Word report of extensive production testing (advertiser app,
management every tab, and the FuseTwo marketing site).
    python scripts/build_prod_extensive_docx.py
Output: docs/FuseTwo_Prod_Extensive_Testing.docx
"""
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT=os.path.join(ROOT,"docs","FuseTwo_Prod_Extensive_Testing.docx")
NAVY=RGBColor(0x1F,0x38,0x64); GREY=RGBColor(0x60,0x60,0x60); RED=RGBColor(0x9C,0x00,0x06); GREEN=RGBColor(0x00,0x61,0x00); AMBER=RGBColor(0x9C,0x65,0x00)

def shade(cell,fill):
    shd=OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:fill"),fill); cell._tc.get_or_add_tcPr().append(shd)
def setc(cell,text,bold=False,color=None,size=9):
    cell.text=""; r=cell.paragraphs[0].add_run(text); r.font.size=Pt(size); r.font.bold=bold
    if color: r.font.color.rgb=color
def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Light Grid Accent 1"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers): setc(t.rows[0].cells[i],h,bold=True,color=RGBColor(0xFF,0xFF,0xFF)); shade(t.rows[0].cells[i],"2E5C8A")
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):
            txt=v[0] if isinstance(v,tuple) else v; col=v[1] if isinstance(v,tuple) else None
            setc(c[i],txt,color=col)
    for i,w in enumerate(widths):
        for row in t.rows: row.cells[i].width=Inches(w)
def head(doc,t): h=doc.add_heading(level=1); h.add_run(t).font.color.rgb=NAVY
def body(doc,t,italic=False,color=None,size=10.5):
    p=doc.add_paragraph(); r=p.add_run(t); r.font.size=Pt(size); r.italic=italic
    if color: r.font.color.rgb=color
def pic(doc,rel,cap):
    path=os.path.join(ROOT,rel)
    if not os.path.exists(path): return
    c=doc.add_paragraph(); r=c.add_run(cap); r.font.size=Pt(9); r.italic=True; r.font.color.rgb=GREY
    doc.add_picture(path,width=Inches(6.2)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER

def main():
    doc=Document(); doc.styles["Normal"].font.name="Calibri"; doc.styles["Normal"].font.size=Pt(10.5)
    t=doc.add_paragraph(); r=t.add_run("FuseTwo — Extensive Production Testing"); r.font.size=Pt(21); r.font.bold=True; r.font.color.rgb=NAVY
    body(doc,"End-to-end testing on the live production environment: advertiser onboarding, "
             "management portal (every tab + sections), and the FuseTwo marketing site. "
             "Prepared by Syed Shah, QA. 1 September 2026.",italic=True,color=GREY)

    head(doc,"Bugs found on production")
    table(doc,["#","Severity","Bug","Evidence"],[
        ["P-1",("High",RED),"Marketing site (www.fusetwo.com): header “Log In” / “Sign Up” link to the DEV advertiser platform (advertiser.dev.fusetwo.com) on all 9 pages","Live users land on dev, not prod. Regression of #297."],
        ["P-2",("High",RED),"Status-change notification e-mails are not delivered (Send succeeds, nothing arrives)","Inbox has only the signup verify e-mail. Same as dev/staging (B-1)."],
        ["P-3",("Low",AMBER),"E-mail confirmation token is plain base64(email) — no signature/expiry","Same weak token as dev/staging (B-5)."],
    ],[0.5,1.0,4.1,2.0])

    head(doc,"Advertiser onboarding (production)")
    table(doc,["Scenario","Result"],[
        ["Register a new advertiser (SYED TEST QA #1062627)",("PASS",GREEN)],
        ["Verification e-mail received + link resolves",("PASS",GREEN)],
        ["Unverified login → pending-account gate",("PASS",GREEN)],
        ["Change status (Pending → Collect Payment)",("PASS",GREEN)],
        ["Status notification e-mail delivered",("FAIL — P-2",RED)],
    ],[5.0,2.2])

    head(doc,"Management portal (production) — every tab & section")
    body(doc,"Logged in via Windows SSO (shown as “Syed Shah”). Every advertiser-detail tab and "
             "management section rendered with no errors, failed APIs, or error banners.")
    table(doc,["Area","Result"],[
        ["Advertiser Overview + search",("clean",GREEN)],
        ["Detail tabs: Details, Users, Account Summary, Programs, Recent Changes",("clean",GREEN)],
        ["Detail tabs: Managed Account, Product Feeds, Pricing Schedule, Tracking",("clean",GREEN)],
        ["Sections: Revenue, Receivables, Advertiser Balance, Promo Codes",("clean",GREEN)],
        ["CRUD: Edit Advertiser Info (website) — persists",("PASS",GREEN)],
        ["CRUD: Edit Contact Info (city) — persists",("PASS",GREEN)],
        ["CRUD: Add internal note",("PASS",GREEN)],
    ],[5.0,2.2])

    head(doc,"FuseTwo marketing site (www.fusetwo.com)")
    table(doc,["Check","Result"],[
        ["9 pages load (Home, Platform, Brands, Services, Industries, Publishers, About, Privacy, Terms)",("PASS — all HTTP 200",GREEN)],
        ["Responsive: no horizontal overflow at 1440 / 768 / 375",("PASS",GREEN)],
        ["Header Log In / Sign Up point to production",("FAIL — point to DEV (P-1)",RED)],
    ],[5.0,2.2])
    pic(doc,"docs/prod-test/site/site-header.png","Production marketing-site header — Log In / Sign Up link to advertiser.dev.fusetwo.com (P-1)")

    head(doc,"Overall verdict")
    for line in [
        "Production is functionally healthy: onboarding, management (every tab), CRUD and the "
        "marketing pages all work, and the site is responsive.",
        "Two high-severity production issues: the marketing site's Log In / Sign Up point to the "
        "DEV platform (P-1), and status-change notification e-mails are not delivered (P-2).",
        "Production behaves the same as staging — same flows, same two shared bugs (P-2/B-1 and "
        "P-3/B-5). No prod-only functional regressions beyond the dev-link header (P-1).",
    ]:
        doc.add_paragraph(style="List Bullet").add_run(line).font.size=Pt(10.5)

    body(doc,"Test data created on production (please decline/remove): SYED TEST QA mtit7q2s / "
             "#1062627 / syed.shah+prodtestmtit7q2s@flexoffers.com — status now Collect Payment.",
             italic=True,color=AMBER,size=9.5)

    doc.save(OUT); print("wrote",OUT)

if __name__=="__main__":
    main()
