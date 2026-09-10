#!/usr/bin/env python
"""Build a Word doc for the management detail-route bounce bug (with screenshots
and a reference to the recorded video).

    python scripts/build_route_bug_docx.py
Output: docs/FuseTwo_Bug_Management_Route.docx
"""
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "FuseTwo_Bug_Management_Route.docx")
NAVY = RGBColor(0x1F, 0x38, 0x64); GREY = RGBColor(0x60, 0x60, 0x60); RED = RGBColor(0x9C, 0x00, 0x06)


def pic(doc, rel, caption):
    path = os.path.join(ROOT, rel)
    if not os.path.exists(path):
        return
    c = doc.add_paragraph(); r = c.add_run(caption)
    r.font.size = Pt(9); r.italic = True; r.font.color.rgb = GREY
    doc.add_picture(path, width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10.5)

    t = doc.add_paragraph(); r = t.add_run("FuseTwo — Bug: Management advertiser detail deep-link bounces")
    r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = NAVY
    s = doc.add_paragraph(); ss = s.add_run(
        "The old advertiser-detail deep-link no longer opens the record — it silently "
        "redirects to the Advertiser Overview. Reproduced on both dev and staging. "
        "Prepared by Syed Shah, QA. 31 August 2026.")
    ss.font.size = Pt(10.5); ss.italic = True; ss.font.color.rgb = GREY

    def head(txt):
        h = doc.add_heading(level=1); h.add_run(txt).font.color.rgb = NAVY

    meta = doc.add_paragraph()
    m = meta.add_run("Severity: Medium    |    Area: Management portal — advertiser detail routing    "
                     "|    Environment: dev and staging (management.stg.fusetwo.com)")
    m.font.size = Pt(9.5); m.italic = True; m.font.color.rgb = GREY

    head("Steps to reproduce")
    for step in [
        "Open the management portal and sign in (Windows SSO).",
        "In a browser, navigate directly to the old detail URL: "
        "https://management.stg.fusetwo.com/advertiser/detail/1062551 "
        "(any valid advertiser id).",
        "Observe the page that loads.",
        "Now navigate to https://management.stg.fusetwo.com/advertiser/1062551 and compare.",
    ]:
        doc.add_paragraph(style="List Number").add_run(step).font.size = Pt(10.5)

    p = doc.add_paragraph(); p.add_run("Expected: ").bold = True
    p.add_run("The detail deep-link opens the advertiser's detail page (or redirects to the "
              "current detail route). A deep-link should never silently drop the id.").font.size = Pt(10.5)
    p = doc.add_paragraph(); rr = p.add_run("Actual: "); rr.bold = True; rr.font.color.rgb = RED
    p.add_run("/advertiser/detail/{id} redirects to /advertiser/overview — the advertiser is "
              "not opened. The working route is now /advertiser/{id}. There is no redirect from "
              "the old path to the new one, so any saved link, bookmark or automation that used "
              "/advertiser/detail/{id} silently lands on the wrong page.").font.size = Pt(10.5)

    head("Evidence")
    for line in [
        "/advertiser/detail/1062551  ->  /advertiser/overview   (BOUNCED — detail not open)",
        "/advertiser/1062551         ->  /advertiser/1062551     (opens — Actions button visible)",
        "The search-result row links the advertiser id to /advertiser/{id}, confirming the new "
        "route; the old /advertiser/detail/{id} is no longer served.",
        "Reproduced 5/5 times on dev and again on staging.",
    ]:
        doc.add_paragraph(style="List Bullet").add_run(line).font.size = Pt(10.5)

    head("Screenshots")
    pic(doc, "docs/bug-screens/route-old-bounced-stg.png",
        "Old route /advertiser/detail/1062551 → bounced to Advertiser Overview")
    pic(doc, "docs/bug-screens/route-new-opened-stg.png",
        "New route /advertiser/1062551 → advertiser detail opens correctly")

    head("Screen recording")
    p = doc.add_paragraph()
    p.add_run("A screen recording of the reproduction (staging) is saved at:\n").font.size = Pt(10.5)
    r = p.add_run("docs/bug-video/mgmt-detail-route-bounce-staging.webm")
    r.font.name = "Consolas"; r.font.size = Pt(10)
    doc.add_paragraph().add_run(
        "The .webm plays in Chrome/Edge/VLC (double-click, or drag into a browser tab). "
        "It shows the old route bouncing to Overview, then the new route opening the detail."
    ).font.size = Pt(9.5)

    head("Impact")
    for line in [
        "Saved links / bookmarks to /advertiser/detail/{id} no longer work.",
        "Any automation or integration using the old route silently lands on Overview.",
        "Fix: add a server/SPA redirect from /advertiser/detail/{id} to /advertiser/{id}.",
    ]:
        doc.add_paragraph(style="List Bullet").add_run(line).font.size = Pt(10.5)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
