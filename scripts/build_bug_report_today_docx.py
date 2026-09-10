#!/usr/bin/env python
"""
Build a downloadable Word (.docx) report of ONLY today's bugs (28 Aug 2026),
with steps to reproduce and screenshots.

    python scripts/build_bug_report_today_docx.py

Output: docs/FuseTwo_Bug_Report_Today.docx
"""

import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "FuseTwo_Bug_Report_Today.docx")

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x60, 0x60, 0x60)
RED = RGBColor(0x9C, 0x00, 0x06)

DATE = "28 August 2026"

# Bugs found TODAY (this session's app crawl). id, severity, title, area, env,
# steps[], expected, actual, evidence, affected[], screenshots[]
BUGS = [
    (
        "BUG-1", "High",
        "Programs API returns HTTP 500 across the whole app",
        "Advertiser platform — Programs API (GET /api/{advertiserId}/programs)",
        "dev  (advertiser 1062551 / playwright@fusetwo.com)",
        [
            "Sign in to the advertiser platform (advertiser.dev.fusetwo.com).",
            "Open the Programs page (or Dashboard, Campaigns, Creatives, Coupons & Offers, "
            "Products, Reports > Sales Detailed, Reports > Clicks Detailed).",
            "Open the browser DevTools > Network (or Console) and watch the request to "
            "/api/1062551/programs.",
        ],
        "GET /api/{advertiserId}/programs returns HTTP 200 with the program list (or an "
        "empty list for a new account), and the page renders without an error.",
        "GET /api/1062551/programs returns HTTP 500 with body "
        '{"result":null,"errorMessage":"Oops! Something went wrong. Please try again in a '
        'few moments."}. It fails on every page that calls it, and the browser console '
        "logs an error each time. Reproduced consistently on repeated loads (verified "
        "twice).",
        "The request is authenticated (a valid Bearer JWT is sent), so this is a "
        "server-side crash, not an auth failure. The same generic error surfaces to the "
        "user as an 'Oops! Something went wrong' toast (seen on Message Center). "
        "Note: the test account is currently in 'Initial Setup' status (not Approved); "
        "the 500 may correlate with that state, but the API should degrade gracefully "
        "(empty list / proper status), not return a 500.",
        [
            "/app/dashboard",
            "/app/programs",
            "/app/campaigns",
            "/app/creatives",
            "/app/couponsandoffers",
            "/app/products",
            "/app/reports/salesdetailed",
            "/app/reports/clicksdetailed",
        ],
        ["docs/bug-screens/programs.png", "docs/bug-screens/messagecenter.png"],
    ),
]


def shade_header(cell, fill="2E5C8A"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def main():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    r = title.add_run("FuseTwo — Bug Report (Today)")
    r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = NAVY
    sub = doc.add_paragraph()
    s = sub.add_run(
        f"Bugs found on {DATE} during QA automation of the advertiser platform. "
        "Each has steps to reproduce, expected vs actual, evidence and screenshots. "
        "Prepared by Syed Shah, QA. Environment: dev."
    )
    s.font.size = Pt(10.5); s.italic = True; s.font.color.rgb = GREY

    # summary
    doc.add_paragraph()
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    for i, h in enumerate(["ID", "Severity", "Title", "Area"]):
        c = tbl.rows[0].cells[i]; c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.font.bold = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_header(c)
    for bug in BUGS:
        cells = tbl.add_row().cells
        for i, v in enumerate([bug[0], bug[1], bug[2], bug[3]]):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(v); run.font.size = Pt(9)
            if i == 1:
                run.font.color.rgb = RED; run.font.bold = True

    for bug in BUGS:
        bug_id, sev, title_t, area, envr, steps, expected, actual, evidence, affected, shots = bug
        doc.add_paragraph()
        h = doc.add_heading(level=1)
        h.add_run(f"{bug_id}: {title_t}").font.color.rgb = NAVY

        meta = doc.add_paragraph()
        m = meta.add_run(f"Severity: {sev}    |    Area: {area}    |    Environment: {envr}")
        m.font.size = Pt(9.5); m.italic = True; m.font.color.rgb = GREY

        doc.add_paragraph().add_run("Steps to reproduce").bold = True
        for step in steps:
            doc.add_paragraph(style="List Number").add_run(step).font.size = Pt(10.5)

        p = doc.add_paragraph(); p.add_run("Expected: ").bold = True; p.add_run(expected).font.size = Pt(10.5)
        p = doc.add_paragraph(); rr = p.add_run("Actual: "); rr.bold = True; rr.font.color.rgb = RED
        p.add_run(actual).font.size = Pt(10.5)
        p = doc.add_paragraph(); p.add_run("Evidence: ").bold = True; p.add_run(evidence).font.size = Pt(10.5)

        doc.add_paragraph().add_run("Affected pages").bold = True
        for a in affected:
            doc.add_paragraph(style="List Bullet").add_run(a).font.size = Pt(10.5)

        for shot in shots:
            path = os.path.join(ROOT, shot)
            if os.path.exists(path):
                cap = doc.add_paragraph(); cr = cap.add_run(f"Screenshot: {os.path.basename(shot)}")
                cr.font.size = Pt(9); cr.italic = True; cr.font.color.rgb = GREY
                doc.add_picture(path, width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
